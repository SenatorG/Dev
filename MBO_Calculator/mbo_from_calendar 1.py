import os
import math
from datetime import datetime, date
import pandas as pd
try:
    import win32com.client
except ImportError:
    raise SystemExit(
        "pywin32 is not installed.\n\n"
        "Run:\n\n"
        "    pip install pywin32\n"
    )
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, PatternFill, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# ---------- CONFIG ----------
CALENDAR_FILE = "CalendarExport.xlsx"  # single Excel file

# ---------- HELPER FUNCTIONS ----------
def round_up_to_half_hour(minutes: float) -> int:
    """Round minutes up to the next multiple of 30 (30, 60, 90, ...)."""
    if minutes <= 0:
        return 0
    return int(math.ceil(minutes / 30.0) * 30)

def classify_mbo(outlook_cat, title, location) -> str:
    """
    Return one of:
    - 'Enablement'
    - 'Opportunity Support'
    - 'GTO/CTO Activities'
    based on Outlook category + title + location.
    Safely handles NaN / non-string values from pandas.
    """

    def norm(val):
        if pd.isna(val):
            return ""
        return str(val).strip()

    oc = norm(outlook_cat)
    t = norm(title).lower()
    loc = norm(location).lower()

    # 1) Straightforward mappings by Outlook category
    if oc in ["OCTO Projects", "Product Group Meeting", "Leadership"]:
        return "GTO/CTO Activities"
    if oc in ["Mentorship", "Internal Admin"]:
        return "Enablement"
    if oc in ["Customer Meeting", "EBC", "Field Engagement"]:
        return "Opportunity Support"
    if oc == "Company Travel":
        return "Opportunity Support"

    # 2) NATO AI Activity
    if oc == "NATO AI Activity":
        gto_keywords = [
            "steerco", "steer co", "steering", "council",
            "core team", "governance", "field ai council",
            "ambassadors", "ctoa", "data platform",
        ]
        if any(k in t for k in gto_keywords):
            return "GTO/CTO Activities"

        opp_keywords = [
            "launch", "prep", "readiness", "review and next steps",
            "demo", "account", "intro",
        ]
        if any(k in t for k in opp_keywords):
            return "Opportunity Support"
        return "Enablement"

    # 3) Partner Meeting
    if oc == "Partner Meeting":
        enablement_keywords = [
            "bootcamp", "training", "office hours", "sync call",
            "weekly touchpoint", "touchpoint", "bi-weekly", "update",
        ]
        if any(k in t for k in enablement_keywords):
            return "Enablement"

        opp_keywords = [
            "appworld", "strategy", "qbr", "dinner", "happy hour", "event",
        ]
        if any(k in t for k in opp_keywords):
            return "Opportunity Support"
        return "Opportunity Support"

    # 4) Categories that usually don’t count: fall through to context
    if oc in [
        "Holiday", "Optional", "Personal",
        "Tracked To Dynamics 365",
        "Tracked To Dynamics 365 (Undeliverable)",
    ]:
        pass

    # 5) No / unknown category: infer from title + location
    gto_keywords = [
        "steerco", "steering", "council", "governance", "core team",
        "automation platform", "ctoa", "community of practice", "genai community",
        "office hours - na & latam", "dap ambassadors",
    ]
    if any(k in t for k in gto_keywords):
        return "GTO/CTO Activities"

    enablement_keywords = [
        "bootcamp", "training", "dojo", "education", "enablement",
        "office hours", "community call", "success circle", "overview of expectations",
        "welcome dinner", "breakout:", "ai solutions training", "ask me anything ai",
        "ai dojo", "review blog", "social post",
    ]
    if any(k in t for k in enablement_keywords):
        return "Enablement"

    opp_keywords = [
        "account", "customer", "strategy", "qbr", "prep call",
        "intro", "dinner", "lunch", "happy hour", "conference",
        "appworld", "cummins", "bell", "f5", "shi", "cdw",
    ]
    if any(k in t for k in opp_keywords) or any(
        k in loc for k in ["vegas", "tempe"]
    ):
        return "Opportunity Support"

    # Fallback
    return "Enablement"

def parse_yyyy_mm_dd(s: str) -> date:
    s = (s or "").strip()
    return datetime.strptime(s, "%Y-%m-%d").date()

def fetch_outlook_calendar(week_start: date, week_end: date) -> pd.DataFrame:
    """
    Pull calendar items from default Outlook calendar for the given date range.
    Returns DataFrame with:
    EventID, Title, Location, OutlookCategories,
    EventDate, StartDateTime, DurationMinutes
    """
    outlook = win32com.client.Dispatch("Outlook.Application")
    ns = outlook.GetNamespace("MAPI")
    calendar = ns.GetDefaultFolder(9)  # 9 = olFolderCalendar
    items = calendar.Items
    items.IncludeRecurrences = True
    items.Sort("[Start]")

    start_str = week_start.strftime("%m/%d/%Y 00:00")
    end_str = week_end.strftime("%m/%d/%Y 23:59")
    restriction = f"[Start] >= '{start_str}' AND [Start] <= '{end_str}'"
    restricted = items.Restrict(restriction)

    rows = []
    for item in restricted:
        try:
            subject = item.Subject
            location = item.Location or ""
            categories = item.Categories or ""
            start = item.Start  # COM -> Python datetime
            duration = item.Duration or 0
            event_id = getattr(item, "GlobalAppointmentID", None) or item.EntryID
        except Exception:
            continue

        # Exclude Personal entries
        if "personal" in categories.lower():
            continue

        # Normalize remote location:
        loc_lower = location.lower()
        if (
            "microsoft teams meeting" in loc_lower
            or "https://" in loc_lower
            or "see body of invite for location/video conference details" in loc_lower
        ):
            location = "Remote"

        event_date_str = start.date().strftime("%Y-%m-%d")
        start_dt_str = start.strftime("%Y-%m-%d %H:%M")

        rows.append(
            {
                "Title": subject,
                "Location": location,
                "OutlookCategories": categories,
                "EventDate": event_date_str,
                "StartDateTime": start_dt_str,
                "DurationMinutes": duration,
                "EventID": event_id,
            }
        )

    return pd.DataFrame(rows)

def apply_mbo_logic(df: pd.DataFrame) -> pd.DataFrame:
    """
    Compute MBO_Category, RoundedMinutes, PSS_Hours for every row.
    Ensure EnteredInOtherSystem column exists (preserve existing values).
    """
    if df.empty:
        return df

    df = df.copy()
    df["DurationMinutes"] = pd.to_numeric(
        df["DurationMinutes"], errors="coerce"
    ).fillna(0)

    df["RoundedMinutes"] = df["DurationMinutes"].apply(round_up_to_half_hour)
    df["PSS_Hours"] = df["RoundedMinutes"] / 60.0

    # Always recompute MBO_Category for all rows
    df["MBO_Category"] = df.apply(
        lambda row: classify_mbo(
            row.get("OutlookCategories", ""),
            row.get("Title", ""),
            row.get("Location", ""),
        ),
        axis=1,
    )

    # Ensure tracking column exists, but do not overwrite existing values
    if "EnteredInOtherSystem" not in df.columns:
        df["EnteredInOtherSystem"] = ""

    return df

def sort_mbo(df: pd.DataFrame) -> pd.DataFrame:
    """
    Sort by MBO category (Enablement → GTO/CTO Activities → Opportunity Support),
    then by EventDate ascending.
    """
    if df.empty:
        return df

    df = df.copy()
    category_order = {
        "Enablement": 0,
        "GTO/CTO Activities": 1,
        "Opportunity Support": 2,
    }
    df["__cat_order"] = df["MBO_Category"].map(category_order).fillna(999)
    df["__event_dt"] = pd.to_datetime(df["EventDate"], errors="coerce")

    df = df.sort_values(
        by=["__cat_order", "__event_dt", "Title"],
        ascending=[True, True, True],
        kind="mergesort",
    )

    df = df.drop(columns=["__cat_order", "__event_dt"])
    return df

def apply_excel_formatting(path: str) -> None:
    """
    Post-process the Excel file to:
    - Bold + fill header row
    - Add Yes/No dropdown to EnteredInOtherSystem
    - Highlight entire row green when EnteredInOtherSystem = Yes
    (EventID is at the far right and not wrapped)
    """
    wb = load_workbook(path)
    ws = wb.active

    max_row = ws.max_row
    max_col = ws.max_column

    # Header formatting: bold + light gray fill
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_font = Font(bold=True)
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill

    # Map headers to column indices
    header_row = 1
    col_index = {}
    for col in range(1, max_col + 1):
        header = ws.cell(row=header_row, column=col).value
        if isinstance(header, str):
            col_index[header] = col

    # EventID: explicitly ensure no wrapping
    if "EventID" in col_index:
        ev_col = col_index["EventID"]
        for row in range(2, max_row + 1):
            cell = ws.cell(row=row, column=ev_col)
            cell.alignment = Alignment(
                wrap_text=False,
                horizontal=cell.alignment.horizontal,
                vertical=cell.alignment.vertical,
            )

    # Data validation dropdown for EnteredInOtherSystem
    if "EnteredInOtherSystem" in col_index:
        ent_col = col_index["EnteredInOtherSystem"]
        ent_col_letter = get_column_letter(ent_col)

        dv = DataValidation(
            type="list",
            formula1='"Yes,No"',
            allow_blank=True,
        )
        dv.errorTitle = "Invalid value"
        dv.error = "Please select Yes or No."
        dv.showErrorMessage = True

        ws.add_data_validation(dv)
        dv_range = f"{ent_col_letter}2:{ent_col_letter}{max_row}"
        dv.add(dv_range)

        # Conditional formatting: entire row green when EnteredInOtherSystem = "Yes"
        last_col_letter = get_column_letter(max_col)
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        rule = FormulaRule(
            formula=[f'${ent_col_letter}2="Yes"'],
            fill=green_fill,
        )
        ws.conditional_formatting.add(
            f"A2:{last_col_letter}{max_row}",
            rule,
        )

    wb.save(path)

def append_or_create_calendar_file(df_new: pd.DataFrame, path: str) -> None:
    """
    Append df_new into path (XLSX).
    Create if missing.
    De-duplicate on (EventID, StartDateTime) if both available, else on EventID.
    Recompute MBO logic for all rows.
    Sort by MBO category then EventDate.
    Column order:
    - Left: EnteredInOtherSystem, MBO_Category, Title, Location,
      DurationMinutes, RoundedMinutes, PSS_Hours
    - Middle: any other columns not explicitly placed
    - Right: OutlookCategories, EventDate, StartDateTime, EventID
    """
    if os.path.isfile(path):
        existing = pd.read_excel(path)

        # Ensure tracking column exists; do NOT compute categories yet
        if "EnteredInOtherSystem" not in existing.columns:
            existing["EnteredInOtherSystem"] = ""

        # Align columns
        for col in df_new.columns:
            if col not in existing.columns:
                existing[col] = None
        for col in existing.columns:
            if col not in df_new.columns:
                df_new[col] = None

        combined = pd.concat([existing, df_new], ignore_index=True)
    else:
        combined = df_new.copy()

    # Remove any legacy WeekStart/WeekEnd if present
    for col in ["WeekStart", "WeekEnd"]:
        if col in combined.columns:
            combined = combined.drop(columns=[col])

    # Ensure tracking column exists before recomputing logic
    if "EnteredInOtherSystem" not in combined.columns:
        combined["EnteredInOtherSystem"] = ""

    # De-duplicate
    if "EventID" in combined.columns and "StartDateTime" in combined.columns:
        combined = combined.drop_duplicates(
            subset=["EventID", "StartDateTime"], keep="last"
        )
    elif "EventID" in combined.columns:
        combined = combined.drop_duplicates(subset=["EventID"], keep="last")

    # Recompute MBO logic for all rows (fixes any bad MBO_Category values)
    combined = apply_mbo_logic(combined)

    # Sort by MBO category + EventDate
    combined = sort_mbo(combined)

    # Desired column order
    core_left = [
        "EnteredInOtherSystem",
        "MBO_Category",
        "Title",
        "Location",
        "DurationMinutes",
        "RoundedMinutes",
        "PSS_Hours",
    ]
    right_last = ["OutlookCategories", "EventDate", "StartDateTime", "EventID"]

    cols = list(combined.columns)
    left_cols = [c for c in core_left if c in cols]
    right_cols = [c for c in right_last if c in cols]
    middle_cols = [c for c in cols if c not in left_cols + right_cols]
    final_cols = left_cols + middle_cols + right_cols
    combined = combined[final_cols]

    combined.to_excel(path, index=False, engine="openpyxl")
    apply_excel_formatting(path)
    print(f"Saved {len(combined)} total events to {path}.")

# ---------- MAIN ----------
def main():
    # 1) Prompt for week
    start_str = input("Enter work week START date (YYYY-MM-DD): ").strip()
    end_str = input("Enter work week END date   (YYYY-MM-DD): ").strip()

    try:
        week_start = parse_yyyy_mm_dd(start_str)
        week_end = parse_yyyy_mm_dd(end_str)
    except ValueError:
        print("Could not parse dates.\nUse YYYY-MM-DD.")
        return

    if week_end < week_start:
        print("End date is before start date.\nExiting.")
        return

    # 2) Fetch events from Outlook
    print("Pulling calendar items from Outlook...")
    df_new = fetch_outlook_calendar(week_start, week_end)
    if df_new.empty:
        print("No Outlook calendar items found in the specified date range.")
        return

    # 3) Append/create single Excel file with everything
    out_path = os.path.join(os.getcwd(), CALENDAR_FILE)
    append_or_create_calendar_file(df_new, out_path)

if __name__ == "__main__":
    main()